import fitz
import re
from pptx import Presentation
from pptx.util import Pt, Inches

def pdf_to_pptx_teaching_style(pdf_path, pptx_path):
    prs = Presentation()

    # Đọc toàn bộ text
    doc = fitz.open(pdf_path)
    text = ""
    for page in doc:
        text += page.get_text()

    # Chuẩn hóa
    text = re.sub(r'\n+', '\n', text)

    # Tách theo "Câu X."
    parts = re.split(r'(Câu\s+\d+\.)', text)

    questions = []
    for i in range(1, len(parts), 2):
        title = parts[i].strip()
        content = parts[i+1].strip()
        questions.append((title, content))

    for title, content in questions:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # slide trắng

        # ===== TIÊU ĐỀ =====
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3),
            Inches(9), Inches(1)
        )
        title_tf = title_box.text_frame
        title_tf.text = title
        title_tf.paragraphs[0].font.size = Pt(34)
        title_tf.paragraphs[0].font.bold = True

        # ===== NỘI DUNG =====
        content_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.4),
            Inches(9), Inches(5)
        )
        tf = content_box.text_frame
        tf.clear()

        lines = content.split("\n")

        for line in lines:
            line = line.strip()
            if not line:
                continue

            p = tf.add_paragraph()
            p.text = line

            # Phân biệt câu hỏi & đáp án
            if re.match(r'[A-D]\.', line):
                p.font.size = Pt(24)
                p.level = 1
            else:
                p.font.size = Pt(26)
                p.level = 0

    prs.save(pptx_path)
    print("✅ ĐÃ TẠO PPT DẠY HỌC – MỖI CÂU 1 SLIDE")

# ====== CHẠY ======
pdf_input = "Made 1202.pdf"
pptx_output = "De_1202_BAN_DEP_DAY_HOC.pptx"

pdf_to_pptx_teaching_style(pdf_input, pptx_output)
